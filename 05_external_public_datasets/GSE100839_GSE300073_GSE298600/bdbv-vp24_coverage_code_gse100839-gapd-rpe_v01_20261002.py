"""Gap-D2: EBOV infection of human retinal pigment epithelium (GSE100839,
PMID 28721309).

Why this dataset matters here
-----------------------------
The manuscript's partition places retinal pigment epithelium at the bottom of the
interferon-tone axis and therefore names it a candidate sanctuary compartment,
on the reasoning that a cell with little interferon tone extracts little penalty
from a VP24 defect. That reasoning has a testable premise: infection of human
RPE should provoke at most a modest interferon-stimulated response. GSE100839
deposits a host transcriptome for exactly that cell type and virus: ARPE-19
cells, mock versus EBOV, 24 hours, triplicate. Its supplementary table names its
columns explicitly (1_24C..3_24C, 1_24E..3_24E), so no column map has to be
assumed.

Frozen readout
--------------
R1  the readout is the manuscript's own 17-gene interferon-stimulated module
    (ISG17), plus its three score components, computed from counts per million
    on the full gene space.
R2  the result is reported as the difference between infected and mock wells in
    log2 CPM, with the per-gene values given so the module is auditable.
R3  this is an immortalised line and a single time point, so the outcome is
    reported as a constraint on the sanctuary premise, never as a test of the
    gradient itself.
"""

from __future__ import annotations

import csv
import json
import math
import os
import statistics as st

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, "raw")
XLSX = os.path.join(RAW, "GSE100839_Read_Counts.xlsx")

ISG17 = ["ISG15", "MX1", "MX2", "OAS1", "OAS2", "IFIT1", "IFIT2", "IFIT3",
         "BST2", "RSAD2", "USP18", "STAT1", "IRF7", "IFI27", "IFI6", "DDX58",
         "IFIH1"]
TYPE_I = ["IFNAR1", "IFNAR2"]
TYPE_III = ["IFNLR1", "IL10RB"]
PRIMING = ["ISG15", "MX1"]
MOCK = ["1_24C", "2_24C", "3_24C"]
EBOV = ["1_24E", "2_24E", "3_24E"]


def load():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [str(x).strip() for x in next(it)]
    cols = header[1:]
    genes = {}
    for row in it:
        g = str(row[0]).strip()
        if not g or g in genes:
            continue
        try:
            genes[g] = [float(x) for x in row[1:]]
        except (TypeError, ValueError):
            continue
    return cols, genes


def main():
    cols, counts = load()
    n = len(cols)
    ci = {c: i for i, c in enumerate(cols)}
    totals = [sum(v[i] for v in counts.values()) for i in range(n)]
    logcpm = {g: [math.log2(v[i] / totals[i] * 1e6 + 1.0) for i in range(n)]
              for g, v in counts.items()}

    def module(genes):
        present = [g for g in genes if g in logcpm]
        vals = [st.fmean(logcpm[g][i] for g in present) for i in range(n)]
        mock = st.fmean(vals[ci[c]] for c in MOCK)
        ebov = st.fmean(vals[ci[c]] for c in EBOV)
        return present, [round(v, 4) for v in vals], round(mock, 4), round(ebov, 4), \
            round(ebov - mock, 4)

    out = {"dataset": "GSE100839", "pmid": "28721309",
           "design": "ARPE-19 (human retinal pigment epithelial line), mock vs EBOV, 24 h, n = 3 + 3",
           "columns": cols, "library_size_min": int(min(totals)),
           "library_size_max": int(max(totals))}
    for name, genes in (("ISG17", ISG17), ("type_I", TYPE_I),
                        ("type_III", TYPE_III), ("priming", PRIMING)):
        present, vals, mock, ebov, diff = module(genes)
        out[name] = {"genes_used": present,
                     "mock_log2cpm": mock, "ebov_log2cpm": ebov,
                     "difference_log2": diff}
    per_gene = {}
    for g in sorted(set(ISG17 + TYPE_I + TYPE_III)):
        if g not in logcpm:
            continue
        mock = st.fmean(logcpm[g][ci[c]] for c in MOCK)
        ebov = st.fmean(logcpm[g][ci[c]] for c in EBOV)
        per_gene[g] = {"mock": round(mock, 3), "ebov": round(ebov, 3),
                       "difference": round(ebov - mock, 3)}
    out["per_gene"] = per_gene
    out["response_is_modest"] = abs(out["ISG17"]["difference_log2"]) < 1.0
    with open(os.path.join(HERE, "gapd_rpe_summary.json"), "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    with open(os.path.join(HERE, "gapd_rpe_per_gene.tsv"), "w", encoding="utf-8",
              newline="") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["gene", "mock_log2cpm", "ebov_log2cpm", "difference_log2"])
        for g, v in per_gene.items():
            w.writerow([g, v["mock"], v["ebov"], v["difference"]])
    print(json.dumps({k: v for k, v in out.items() if k != "per_gene"},
                     ensure_ascii=False, indent=1))
    print("\nper-gene (log2 CPM):")
    for g, v in per_gene.items():
        print(f"  {g:8s} mock {v['mock']:7.2f}  EBOV {v['ebov']:7.2f}  diff {v['difference']:+6.2f}")


if __name__ == "__main__":
    main()
