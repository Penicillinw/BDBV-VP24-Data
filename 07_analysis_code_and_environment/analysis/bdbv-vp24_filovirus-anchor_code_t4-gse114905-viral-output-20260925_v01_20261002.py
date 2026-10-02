"""Test the pre-registered F8/P6 prediction with GSE114905.

GSE114905 (Huh7 cells, 4 orthoebolaviruses, 1-3 dpi) is NOT a host transcriptome:
the supplementary workbooks report the read counts / TPM of the seven VIRAL mRNAs.
That makes it a direct test of the frozen prediction:

    P1/F8: in a high-Delta compartment (Huh7 = hepatocyte lineage, Delta rank 34)
           BDBV per-cell output should be LOWER than EBOV.

Outputs: analysis/t4_gse114905_20260925/viral_mrna_counts.tsv
         analysis/t4_gse114905_20260925/viral_mrna_summary.json
"""

import glob
import json
import os
import re

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIR = os.path.join(ROOT, "analysis", "t4_gse114905_20260925")

VIRUS_OF_FILE = {
    "BDBV": "BDBV",
    "EBOV": "EBOV",
    "SUDV": "SUDV",
    "RESTV": "RESTV",
}


def main():
    records = []
    for path in sorted(glob.glob(os.path.join(DIR, "*.xlsx"))):
        base = os.path.basename(path)
        virus = next(v for v in VIRUS_OF_FILE if f"_{v}_" in base)
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            day = int(re.search(r"D(\d)\s*\(GE\)", ws.title).group(1))
            header = [c for c in next(ws.iter_rows(values_only=True))]
            idx = {name: i for i, name in enumerate(header) if name}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                ref_col = idx.get("Chromosome")
                records.append(
                    {
                        "virus": virus,
                        "day": day,
                        "gene": row[idx["Name"]],
                        "reference": (row[ref_col] if ref_col is not None
                                      and ref_col < len(row) else "NA"),
                        "reads": float(row[idx["Expression value"]]),
                        "tpm": float(row[idx["TPM"]]),
                    }
                )

    with open(os.path.join(DIR, "viral_mrna_counts.tsv"), "w", encoding="utf-8", newline="") as fh:
        fh.write("virus\tday\tgene\treference\treads\ttpm\n")
        for r in records:
            fh.write(f"{r['virus']}\t{r['day']}\t{r['gene']}\t{r['reference']}\t"
                     f"{r['reads']:.0f}\t{r['tpm']:.2f}\n")

    # total viral reads per virus/day = whole-genome expression proxy
    totals = {}
    for r in records:
        totals.setdefault((r["virus"], r["day"]), 0.0)
        totals[(r["virus"], r["day"])] += r["reads"]

    viruses = sorted({v for v, _ in totals})
    days = sorted({d for _, d in totals})
    print("total viral mRNA reads per time point (sum of the 7 viral genes)")
    print("virus " + "".join(f"{'D'+str(d):>12s}" for d in days))
    for v in viruses:
        print(f"{v:5s} " + "".join(f"{totals.get((v,d), float('nan')):>12,.0f}" for d in days))

    print()
    print("ratios at each time point")
    ratios = {}
    for v in viruses:
        if v == "BDBV":
            continue
        for d in days:
            b, o = totals.get(("BDBV", d)), totals.get((v, d))
            if b and o:
                ratios[f"BDBV/{v}_D{d}"] = b / o
                print(f"  BDBV/{v} D{d}: {b/o:.3f}   (BDBV {b:,.0f} vs {v} {o:,.0f})")

    print()
    print("BDBV gene-level share per day (reads / total) vs EBOV")
    for d in days:
        for gene in ["NP", "VP35", "VP40", "GP", "VP30", "VP24", "L"]:
            def share(v):
                num = next((r["reads"] for r in records
                            if r["virus"] == v and r["day"] == d and r["gene"] == gene), None)
                den = totals.get((v, d))
                return None if num is None or not den else num / den
            sb, se = share("BDBV"), share("EBOV")
            if sb and se:
                print(f"  D{d} {gene:4s} BDBV={sb:.4f} EBOV={se:.4f} diff={sb-se:+.4f}")

    json.dump(
        {
            "source": "GSE114905 supplementary workbooks (ftp.ncbi.nlm.nih.gov/geo)",
            "design_note": "seven VIRAL mRNAs per virus per day; NOT a host transcriptome",
            "totals_reads": {f"{v}_D{d}": totals[(v, d)] for v, d in totals},
            "ratios_vs_BDBV": ratios,
            "caveats": [
                "one library per virus per time point (no biological replicates)",
                "inoculum / MOI normalisation not verified from the series metadata",
                "read counts are not normalised to host reads or to viral genome copies",
            ],
        },
        open(os.path.join(DIR, "viral_mrna_summary.json"), "w", encoding="utf-8"),
        ensure_ascii=False,
        indent=1,
    )


if __name__ == "__main__":
    main()
